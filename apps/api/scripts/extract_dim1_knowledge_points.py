from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import fitz
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

DEFAULT_EXCEL_PATH = Path("C:/Users/admin/Desktop/新-【小数】计算分级表.xlsx")
DEFAULT_PDF_PATH = Path("C:/Users/admin/Desktop/2024知识树.pdf")
DEFAULT_OUTPUT_PATH = (
    ROOT / "app" / "services" / "parser" / "dim1_knowledge_points.json"
)

SCHOOL_SHEET_NAME = "分级表-7月1日更新"
SCHOOL_BRANCH = "校内计算"

LEVEL_TO_GRADE = {
    "L1": 1,
    "L2": 1,
    "L3": 2,
    "L4": 2,
    "L5": 3,
    "L6": 3,
    "L7": 4,
    "L8": 4,
    "L9": 5,
    "L10": 5,
    "L11": 6,
    "L12": 6,
}

LECTURE_RE = re.compile(r"^第(\d+)讲-(\d+)$")
GRADE_RE = re.compile(r"^[1-6]年级$")


@dataclass(frozen=True)
class TextLine:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2


@dataclass(frozen=True)
class BranchRegion:
    branch: str
    x0: float
    y0: float
    x1: float
    y1: float

    def contains(self, line: TextLine) -> bool:
        return self.x0 <= line.cx <= self.x1 and self.y0 <= line.cy <= self.y1


@dataclass(frozen=True)
class LectureBlock:
    branch: str
    grade: int
    lecture_label: str
    lecture_sequence: int
    title: str
    subtopics: tuple[str, ...]


# These are the visual branch blocks on the one-page 2024 knowledge tree.
# They intentionally bound extraction by branch layout instead of keywords.
GAOSI_BRANCH_REGIONS = (
    BranchRegion("计算", 1470, 185, 1925, 1210),
)

GAOSI_GRADE_RANGES = {
    "计算": ((range(1, 6), 3), (range(6, 9), 4), (range(9, 12), 5), (range(12, 15), 6)),
}

PDF_NOISE_TOKENS = {
    "",
    "知识树",
    "数字谜",
    "数论",
    "计算",
    "数",
    "字",
    "谜",
    "论",
    "t",
}


def _normalize_text(value: object) -> str:
    text = str(value or "").strip()
    text = text.replace("\r", "\n")
    text = re.sub(r"\s*\n\s*", "", text)
    text = re.sub(r"\s+", "", text)
    return text.strip()


def _is_noise_token(text: str) -> bool:
    if text in PDF_NOISE_TOKENS:
        return True
    if GRADE_RE.fullmatch(text):
        return True
    return False


def _is_valid_knowledge_point(text: str) -> bool:
    if not text:
        return False
    if text in PDF_NOISE_TOKENS:
        return False
    if GRADE_RE.fullmatch(text):
        return False
    if LECTURE_RE.fullmatch(text):
        return False
    return True


def _entry(
    *,
    entry_id: str,
    source: str,
    display_name: str,
    track: str,
    grade: int,
    level_code: str,
    branch: str,
    knowledge_point: str,
    parent_topic: str,
    source_file: str,
    source_location: str,
) -> dict[str, object]:
    return {
        "id": entry_id,
        "source": source,
        "display_name": display_name,
        "track": track,
        "grade": grade,
        "level_code": level_code,
        "branch": branch,
        "knowledge_point": knowledge_point,
        "parent_topic": parent_topic,
        "source_file": source_file,
        "source_location": source_location,
    }


def _without_ids(entries: Iterable[dict[str, object]]) -> list[dict[str, object]]:
    stripped: list[dict[str, object]] = []
    for entry in entries:
        next_entry = dict(entry)
        next_entry.pop("id", None)
        stripped.append(next_entry)
    return stripped


def _dedupe_and_assign_ids(entries: Iterable[dict[str, object]]) -> list[dict[str, object]]:
    deduped: list[dict[str, object]] = []
    seen_display_names: set[str] = set()
    counters = {"school": 0, "gaosi": 0}

    for entry in _without_ids(entries):
        display_name = str(entry["display_name"])
        knowledge_point = str(entry["knowledge_point"])
        if display_name in seen_display_names or not _is_valid_knowledge_point(knowledge_point):
            continue
        seen_display_names.add(display_name)

        prefix = "school" if entry["source"] == "school_excel" else "gaosi"
        counters[prefix] += 1
        entry_id = f"dim1_{prefix}_{counters[prefix]:04d}"
        deduped.append(_entry(entry_id=entry_id, **entry))

    return deduped


def extract_school_entries(excel_path: Path) -> list[dict[str, object]]:
    if not excel_path.exists():
        raise FileNotFoundError(excel_path)

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if SCHOOL_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"worksheet not found: {SCHOOL_SHEET_NAME}")

        worksheet = workbook[SCHOOL_SHEET_NAME]
        entries: list[dict[str, object]] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            level_code = _normalize_text(row[0] if row else "")
            if level_code not in LEVEL_TO_GRADE:
                continue

            grade = LEVEL_TO_GRADE[level_code]
            for column_index, value in enumerate(row[1:], start=2):
                knowledge_point = _normalize_text(value)
                if not _is_valid_knowledge_point(knowledge_point):
                    continue

                entries.append(
                    _entry(
                        entry_id="",
                        source="school_excel",
                        display_name=f"校内-{grade}年级-{knowledge_point}",
                        track="校内",
                        grade=grade,
                        level_code=level_code,
                        branch=SCHOOL_BRANCH,
                        knowledge_point=knowledge_point,
                        parent_topic="",
                        source_file=excel_path.name,
                        source_location=(
                            f"Excel sheet={SCHOOL_SHEET_NAME}; row={row_index}; "
                            f"column={column_index}; level={level_code}"
                        ),
                    )
                )
    finally:
        workbook.close()

    return entries


def _extract_pdf_lines(pdf_path: Path) -> list[TextLine]:
    if not pdf_path.exists():
        raise FileNotFoundError(pdf_path)

    lines: list[TextLine] = []
    with fitz.open(str(pdf_path)) as document:
        if document.page_count < 1:
            return lines
        page = document.load_page(0)
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                text = _normalize_text(
                    "".join(span.get("text", "") for span in line.get("spans", []))
                )
                if not text:
                    continue
                x0, y0, x1, y1 = line["bbox"]
                lines.append(TextLine(text, x0, y0, x1, y1))
    return lines


def _gaosi_grade_for_lecture(branch: str, lecture_sequence: int) -> int:
    for sequence_range, grade in GAOSI_GRADE_RANGES[branch]:
        if lecture_sequence in sequence_range:
            return grade
    raise ValueError(f"cannot map {branch} lecture sequence {lecture_sequence} to grade")


def _find_title_line(lines: list[TextLine], lecture_line: TextLine) -> TextLine | None:
    for max_y_delta, max_x_delta in ((32, 75), (45, 120)):
        candidates = [
            candidate
            for candidate in lines
            if candidate.cy > lecture_line.cy
            and candidate.cy - lecture_line.cy < max_y_delta
            and abs(candidate.cx - lecture_line.cx) < max_x_delta
            and not LECTURE_RE.fullmatch(candidate.text)
            and _is_valid_knowledge_point(candidate.text)
        ]
        if candidates:
            return min(
                candidates,
                key=lambda candidate: (
                    candidate.cy - lecture_line.cy,
                    abs(candidate.cx - lecture_line.cx),
                ),
            )
    return None


def _parse_branch_lectures(region: BranchRegion, all_lines: list[TextLine]) -> list[LectureBlock]:
    branch_lines = [
        line
        for line in all_lines
        if region.contains(line) and not _is_noise_token(line.text)
    ]
    branch_lines.sort(key=lambda line: (line.cy, line.cx))

    lecture_infos: list[tuple[TextLine, TextLine, int, int]] = []
    for line in branch_lines:
        match = LECTURE_RE.fullmatch(line.text)
        if not match:
            continue
        lecture_sequence = int(match.group(2))
        title_line = _find_title_line(branch_lines, line)
        if title_line is None:
            continue
        grade = _gaosi_grade_for_lecture(region.branch, lecture_sequence)
        lecture_infos.append((line, title_line, lecture_sequence, grade))

    title_line_ids = {id(title_line) for _, title_line, _, _ in lecture_infos}
    subtopics_by_lecture_id: dict[int, list[str]] = {id(line): [] for line, _, _, _ in lecture_infos}

    for line in branch_lines:
        if LECTURE_RE.fullmatch(line.text) or id(line) in title_line_ids:
            continue
        if not _is_valid_knowledge_point(line.text):
            continue

        candidates: list[tuple[float, float, TextLine]] = []
        for lecture_line, title_line, _, _ in lecture_infos:
            if line.cy <= title_line.cy + 1:
                continue
            x_delta = abs(line.cx - lecture_line.cx)
            if x_delta <= 115:
                candidates.append((title_line.cy, x_delta, lecture_line))

        if not candidates:
            continue

        candidates.sort(key=lambda item: (item[0], -item[1]), reverse=True)
        subtopics_by_lecture_id[id(candidates[0][2])].append(line.text)

    blocks: list[LectureBlock] = []
    for lecture_line, title_line, lecture_sequence, grade in lecture_infos:
        subtopics: list[str] = []
        seen_subtopics: set[str] = set()
        for subtopic in subtopics_by_lecture_id[id(lecture_line)]:
            if subtopic in seen_subtopics:
                continue
            seen_subtopics.add(subtopic)
            subtopics.append(subtopic)

        blocks.append(
            LectureBlock(
                branch=region.branch,
                grade=grade,
                lecture_label=lecture_line.text,
                lecture_sequence=lecture_sequence,
                title=title_line.text,
                subtopics=tuple(subtopics),
            )
        )

    return sorted(blocks, key=lambda block: block.lecture_sequence)


def extract_gaosi_entries(pdf_path: Path) -> list[dict[str, object]]:
    lines = _extract_pdf_lines(pdf_path)
    entries: list[dict[str, object]] = []

    for region in GAOSI_BRANCH_REGIONS:
        for block in _parse_branch_lectures(region, lines):
            lecture_ref = f"{block.lecture_label} {block.title}"
            for item_type, knowledge_point in (
                ("lecture_title", block.title),
                *[("subtopic", subtopic) for subtopic in block.subtopics],
            ):
                if not _is_valid_knowledge_point(knowledge_point):
                    continue
                entries.append(
                    _entry(
                        entry_id="",
                        source="gaosi_knowledge_tree_pdf",
                        display_name=f"高思导引-{block.grade}年级-{knowledge_point}",
                        track="高思导引",
                        grade=block.grade,
                        level_code="",
                        branch=block.branch,
                        knowledge_point=knowledge_point,
                        parent_topic=block.title,
                        source_file=pdf_path.name,
                        source_location=(
                            f"PDF page=1; branch={block.branch}; grade={block.grade}; "
                            f"lecture={lecture_ref}; item={item_type}"
                        ),
                    )
                )

    return entries


def build_entries(excel_path: Path, pdf_path: Path) -> list[dict[str, object]]:
    return _dedupe_and_assign_ids(
        [
            *extract_school_entries(excel_path),
            *extract_gaosi_entries(pdf_path),
        ]
    )


def write_json(entries: list[dict[str, object]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(entries, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract calculation-dimension knowledge points from school Excel and GaoSi knowledge-tree PDF.",
    )
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL_PATH)
    parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    entries = build_entries(args.excel, args.pdf)
    write_json(entries, args.output)

    by_source: dict[str, int] = {}
    by_branch: dict[str, int] = {}
    for entry in entries:
        by_source[str(entry["source"])] = by_source.get(str(entry["source"]), 0) + 1
        by_branch[str(entry["branch"])] = by_branch.get(str(entry["branch"]), 0) + 1

    print(f"wrote {len(entries)} entries to {args.output}")
    for source, count in sorted(by_source.items()):
        print(f"{source}: {count}")
    for branch, count in sorted(by_branch.items()):
        print(f"{branch}: {count}")


if __name__ == "__main__":
    main()
